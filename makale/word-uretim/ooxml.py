# -*- coding: utf-8 -*-
"""Yerlesik Word grafik (c:chart) ve sekil (wpg/wps) OOXML uretici."""
from xml.sax.saxutils import escape
import io
from openpyxl import Workbook

C  = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
A  = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R  = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

def _num_cache(vals, fmt='General'):
    pts = ''.join('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, v) for i, v in enumerate(vals))
    return ('<c:numCache><c:formatCode>%s</c:formatCode><c:ptCount val="%d"/>%s</c:numCache>'
            % (fmt, len(vals), pts))

def _str_cache(vals):
    pts = ''.join('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, escape(v)) for i, v in enumerate(vals))
    return '<c:strCache><c:ptCount val="%d"/>%s</c:strCache>' % (len(vals), pts)

def _txpr(size=900, bold=0, color='52514E'):
    return ('<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="%d" b="%d">'
            '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:latin typeface="Times New Roman"/></a:defRPr></a:pPr><a:endParaRPr lang="tr-TR"/>'
            '</a:p></c:txPr>' % (size, bold, color))

def bar_chart(cats, series, bar_dir='col', grouping='clustered',
              val_title=None, cat_title=None, legend=None, gap=60, overlap=None,
              val_max=None, sheet='Sayfa1'):
    """series: [{'name':str,'vals':[...],'color':hex|None,'labels':bool}]"""
    sers = []
    for idx, s in enumerate(series):
        col = chr(ord('B') + idx)
        fill = ('<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % s['color']) if s.get('color') else '<a:noFill/>'
        dlbls = ''
        if s.get('labels'):
            skip = ''.join('<c:dLbl><c:idx val="%d"/><c:delete val="1"/></c:dLbl>' % k
                           for k in s.get('skip', []))
            dlbls = ('<c:dLbls>' + skip + '%s<c:dLblPos val="%s"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
                     '<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/>'
                     '<c:showBubbleSize val="0"/></c:dLbls>'
                     % (_txpr(900, 0, '0B0B0B'), s.get('lblpos', 'outEnd')))
        sers.append(
          '<c:ser><c:idx val="%d"/><c:order val="%d"/>'
          '<c:tx><c:strRef><c:f>%s!$%s$1</c:f><c:strCache><c:ptCount val="1"/>'
          '<c:pt idx="0"><c:v>%s</c:v></c:pt></c:strCache></c:strRef></c:tx>'
          '<c:spPr>%s<a:ln><a:noFill/></a:ln></c:spPr>%s'
          '<c:cat><c:strRef><c:f>%s!$A$2:$A$%d</c:f>%s</c:strRef></c:cat>'
          '<c:val><c:numRef><c:f>%s!$%s$2:$%s$%d</c:f>%s</c:numRef></c:val>'
          '</c:ser>'
          % (idx, idx, sheet, col, escape(s['name']), fill, dlbls,
             sheet, len(cats) + 1, _str_cache(cats),
             sheet, col, col, len(cats) + 1, _num_cache(s['vals'])))

    ov = '<c:overlap val="%d"/>' % overlap if overlap is not None else ''
    catax_title = ''
    valax_title = ''
    if cat_title:
        catax_title = ('<c:title><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="900" b="0">'
                       '<a:solidFill><a:srgbClr val="52514E"/></a:solidFill><a:latin typeface="Times New Roman"/>'
                       '</a:defRPr></a:pPr><a:r><a:t>%s</a:t></a:r></a:p></c:rich></c:tx>'
                       '<c:overlay val="0"/></c:title>' % escape(cat_title))
    if val_title:
        valax_title = ('<c:title><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="900" b="0">'
                       '<a:solidFill><a:srgbClr val="52514E"/></a:solidFill><a:latin typeface="Times New Roman"/>'
                       '</a:defRPr></a:pPr><a:r><a:t>%s</a:t></a:r></a:p></c:rich></c:tx>'
                       '<c:overlay val="0"/></c:title>' % escape(val_title))
    vmax = ('<c:max val="%s"/>' % val_max if val_max is not None else '') + '<c:min val="0"/>'
    leg = ''
    if legend:
        leg = ('<c:legend><c:legendPos val="%s"/><c:overlay val="0"/>%s</c:legend>'
               '<c:plotVisOnly val="1"/>' % (legend, _txpr(900)))
    else:
        leg = '<c:plotVisOnly val="1"/>'

    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
      '<c:chartSpace xmlns:c="%s" xmlns:a="%s" xmlns:r="%s">'
      '<c:roundedCorners val="0"/>'
      '<c:chart><c:autoTitleDeleted val="1"/><c:plotArea><c:layout/>'
      '<c:barChart><c:barDir val="%s"/><c:grouping val="%s"/><c:varyColors val="0"/>'
      '%s<c:gapWidth val="%d"/>%s<c:axId val="111111111"/><c:axId val="222222222"/></c:barChart>'
      '<c:catAx><c:axId val="111111111"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
      '<c:delete val="0"/><c:axPos val="%s"/>%s'
      '<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="DCDCD8"/></a:solidFill></a:ln></c:spPr>'
      '%s<c:crossAx val="222222222"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'
      '<c:valAx><c:axId val="222222222"/><c:scaling><c:orientation val="minMax"/>%s</c:scaling>'
      '<c:delete val="0"/><c:axPos val="%s"/>'
      '<c:majorGridlines><c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="ECECEB"/></a:solidFill>'
      '</a:ln></c:spPr></c:majorGridlines>%s<c:numFmt formatCode="General" sourceLinked="1"/>'
      '<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
      '<c:spPr><a:ln><a:noFill/></a:ln></c:spPr>%s<c:crossAx val="111111111"/></c:valAx>'
      '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>%s'
      '</c:chart><c:spPr><a:solidFill><a:srgbClr val="FCFCFB"/></a:solidFill>'
      '<a:ln><a:noFill/></a:ln></c:spPr>'
      '<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>'
      % (C, A, R, bar_dir, grouping, ''.join(sers), gap, ov,
         'b' if bar_dir == 'col' else 'l', catax_title, _txpr(900),
         vmax, 'l' if bar_dir == 'col' else 'b', valax_title, _txpr(900), leg))

def workbook_bytes(cats, series, sheet='Sayfa1'):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    ws['A1'] = ''
    for j, s in enumerate(series): ws.cell(row=1, column=2 + j, value=s['name'])
    for i, c in enumerate(cats):
        ws.cell(row=2 + i, column=1, value=c)
        for j, s in enumerate(series):
            ws.cell(row=2 + i, column=2 + j, value=s['vals'][i])
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
