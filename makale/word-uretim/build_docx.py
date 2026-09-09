# -*- coding: utf-8 -*-
"""Markdown makaleyi, tablolari YERLESIK Word tablosu olan bir .docx'e cevirir.
Sekiller icin yer tutucu paragraflar birakir; grafik ve sekiller ikinci asamada
yerlesik OOXML nesnesi olarak enjekte edilir."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = '/home/user/github.io/makale/derleme-firin-had-optimizasyonu-humanize.md'
DST = '/home/user/github.io/makale/derleme-firin-had-optimizasyonu.docx'

BODY_FONT, BODY_SIZE = 'Times New Roman', Pt(11)
TBL_SIZE, CAP_SIZE = Pt(8.5), Pt(9)
HDR_FILL = 'E8EDF5'

def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexfill)
    tcPr.append(sh)

def set_cell_borders(cell, color='9AA3AE', sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        el = OxmlElement('w:'+edge)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz))
        el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)
        borders.append(el)
    tcPr.append(borders)

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
def add_runs(par, text, size=BODY_SIZE, bold=False, italic=False, font=BODY_FONT):
    for piece in INLINE.split(text):
        if not piece: continue
        b, i = bold, italic
        t = piece
        if piece.startswith('**') and piece.endswith('**'): t, b = piece[2:-2], True
        elif piece.startswith('*') and piece.endswith('*'): t, i = piece[1:-1], True
        elif piece.startswith('`') and piece.endswith('`'): t = piece[1:-1]
        r = par.add_run(t)
        r.bold, r.italic = b, i
        r.font.name, r.font.size = font, size
        r._element.rPr.rFonts.set(qn('w:eastAsia'), font)

def P(doc, text='', size=BODY_SIZE, align=None, italic=False, bold=False,
      space_after=6, space_before=0, indent=None, hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(space_after), Pt(space_before)
    pf.line_spacing = 1.15
    if align: p.alignment = align
    if indent is not None: pf.left_indent = Cm(indent)
    if hanging is not None: pf.first_line_indent = Cm(-hanging)
    if text: add_runs(p, text, size=size, bold=bold, italic=italic)
    return p

def H(doc, text, level):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = BODY_FONT
    r.font.size = Pt(14 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
    r.bold = True
    r._element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
    return h

def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]

def add_table(doc, rows, widths_cm):
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths_cm[ci])
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            add_runs(par, val, size=TBL_SIZE, bold=(ri == 0))
            set_cell_borders(cell)
            if ri == 0: shade(cell, HDR_FILL)
    # sutun genisliklerini tabloya da yaz (Google Docs uyumu)
    grid = t._tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths_cm):
        gc.set(qn('w:w'), str(int(w * 567)))
    return t

# ------------------------------------------------------------------ ayrıştırma
md = open(SRC, encoding='utf-8').read()
md = md.split('\n---\n---\n\n# EK A')[0]        # Ek A ve Ek B Word surumune alinmaz
lines = md.split('\n')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)
st = doc.styles['Normal']
st.font.name, st.font.size = BODY_FONT, BODY_SIZE
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)

USABLE = 16.0
TBL_W = {  # tablo numarasi -> sutun genislikleri (cm)
 1: [11.0, 5.0],
 2: [2.3, 2.6, 2.2, 3.3, 3.1, 2.5],
 3: [2.6, 3.1, 3.1, 4.0, 3.2],
 4: [3.0, 2.4, 3.4, 3.4, 1.6, 2.2],
 5: [2.5, 2.6, 2.3, 2.5, 2.3, 1.9, 1.9],
 6: [2.0, 2.2, 2.4, 2.4, 2.4, 2.3, 2.3],
 7: [2.8, 3.6, 3.2, 3.4, 3.0],
 8: [3.0, 3.0, 2.2, 4.0, 3.8],
 9: [1.3, 3.4, 5.3, 6.0],
}
i, ntab = 0, 0
in_refs = False
while i < len(lines):
    ln = lines[i]
    s = ln.strip()

    if not s or s == '---':
        i += 1; continue

    if s.startswith('> '):                       # surum notu
        P(doc, s[2:], size=Pt(9), italic=True, space_after=10); i += 1; continue

    if s.startswith('# '):
        p = P(doc, s[2:], size=Pt(16), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              space_after=14, space_before=0)
        i += 1; continue

    if s.startswith('## '):
        title = s[3:]
        in_refs = title.startswith('KAYNAKLAR')
        H(doc, title, 1); i += 1; continue

    if s.startswith('### '):
        H(doc, s[4:], 2); i += 1; continue

    if s.startswith('!['):                       # sekil yer tutucu
        m = re.match(r'!\[Şekil (\d+)\]', s)
        P(doc, '[[SEKIL%s]]' % m.group(1), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        i += 1; continue

    if s.startswith('|'):                        # tablo
        block = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i]); i += 1
        rows = [split_row(b) for b in block if not re.match(r'^\|[\s\-:|]+\|$', b.strip())]
        ntab += 1
        w = TBL_W.get(ntab, [USABLE / len(rows[0])] * len(rows[0]))
        add_table(doc, rows, w)
        P(doc, '', space_after=8)
        continue

    if s.startswith('*Şekil') or s.startswith('*Tablo') or (s.startswith('*Not:') or s.startswith('*Tablo,')):
        P(doc, s.strip('*'), size=CAP_SIZE, italic=True,
          align=WD_ALIGN_PARAGRAPH.CENTER if s.startswith('*Şekil') else None,
          space_after=10); i += 1; continue

    if re.match(r'^Tablo \d+\.', s):             # tablo basligi
        P(doc, s, size=CAP_SIZE, bold=True, space_after=3, space_before=8); i += 1; continue

    if in_refs and re.match(r'^\[\d+\]', s):
        P(doc, s, size=Pt(9.5), indent=0.9, hanging=0.9, space_after=4); i += 1; continue

    P(doc, s, align=WD_ALIGN_PARAGRAPH.JUSTIFY); i += 1

doc.save(DST)
print('yazildi:', DST, '| tablo sayisi:', ntab)
